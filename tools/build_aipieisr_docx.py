from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "docs" / "AIPEIISR_PRD_and_System_Architecture.md"
OUTPUT = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else ROOT / "output" / "AIPEIISR_Product_Requirements_and_System_Architecture.docx"

BLUE = "183B5B"
PALE = "EEF3F7"
GREY = "D9E1E8"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(cell, color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr(); tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = tcBorders.find(qn(tag))
        if el is None: el = OxmlElement(tag); tcBorders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),color)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:' + side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'),'dxa')

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    trPr.append(header)

def configure(doc):
    sec = doc.sections[0]
    sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
    styles = doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); styles['Normal'].font.size=Pt(9.5)
    styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.12
    for name,size,color in [('Title',25,'000000'),('Heading 1',16,'000000'),('Heading 2',12,'000000'),('Heading 3',10.5,'000000')]:
        s=styles[name]; s.font.name='Aptos Display' if name!='Heading 3' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),s.font.name); s._element.rPr.rFonts.set(qn('w:hAnsi'),s.font.name); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(15 if name!='Title' else 0); s.paragraph_format.space_after=Pt(6)
    header=sec.header.paragraphs[0]; header.text='AIPEIISR Product Requirements and System Architecture'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('5A6B7A')
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('AIPEIISR  |  Stakeholder Review Draft  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('5A6B7A')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def add_text(p, text, bold=False, italic=False, color=None, size=None):
    parts=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text)
    for part in parts:
        if not part: continue
        r=p.add_run(part[2:-2] if part.startswith('**') else (part[1:-1] if part.startswith('`') else part))
        r.bold=bold or part.startswith('**'); r.italic=italic
        if part.startswith('`'):
            r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:ascii'),'Consolas'); r.font.size=Pt(8.5)
        if color: r.font.color.rgb=RGBColor.from_string(color)
        if size: r.font.size=Pt(size)

def add_table(doc, rows):
    parsed=[]
    for line in rows:
        cells=[x.strip() for x in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells): continue
        parsed.append(cells)
    if not parsed: return
    cols=max(len(x) for x in parsed); table=doc.add_table(rows=0, cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for i,row in enumerate(parsed):
        cells=table.add_row().cells
        if i == 0: repeat_header(table.rows[-1])
        for j in range(cols):
            c=cells[j]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c); borders(c)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            add_text(p, row[j] if j<len(row) else '', bold=(i==0), color=('FFFFFF' if i==0 else None), size=(8.0 if i else 8.2))
            if i==0: shade(c,BLUE)
            elif i%2==0: shade(c,PALE)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def build():
    doc=Document(); configure(doc)
    lines=SOURCE.read_text(encoding='utf-8').splitlines(); i=0; first_title=True
    while i < len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('```'):
            kind=line[3:].strip(); i+=1; block=[]
            while i<len(lines) and not lines[i].startswith('```'): block.append(lines[i]); i+=1
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.right_indent=Inches(.2); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(7)
            pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F3F6F8'); pPr.append(shd)
            add_text(p, ('Diagram definition (Mermaid)\n' if kind=='mermaid' else '')+'\n'.join(block), color='31495F', size=8.0)
            i+=1; continue
        if line.startswith('|'):
            table=[]
            while i<len(lines) and lines[i].startswith('|'): table.append(lines[i]); i+=1
            add_table(doc,table); continue
        m=re.match(r'^(#{1,3})\s+(.*)$',line)
        if m:
            level=len(m.group(1)); text=m.group(2).strip()
            if level==1:
                p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_text(p,text)
                if first_title:
                    meta=doc.add_paragraph(); meta.paragraph_format.space_after=Pt(14); add_text(meta,'Product Requirements Document and System Architecture  |  Version 1.1  |  Stakeholder Review Draft',color='5A6B7A',size=10); first_title=False
            else:
                p=doc.add_paragraph(style=('Heading 1' if level==2 else 'Heading 2')); add_text(p,text)
            i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); add_text(p,line[2:]); i+=1; continue
        if re.match(r'^\d+\.\s+',line):
            p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); add_text(p,re.sub(r'^\d+\.\s+','',line)); i+=1; continue
        if line.startswith('> '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.right_indent=Inches(.2); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(6); add_text(p,line[2:],italic=True,color='31495F'); i+=1; continue
        p=doc.add_paragraph(); add_text(p,line); i+=1
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.core_properties.title='AIPEIISR Product Architecture and UI Blueprint'; doc.core_properties.author='AIPEIISR'; doc.core_properties.subject='Implementation-ready human-led AI-augmented election information integrity platform'
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__=='__main__': build()
